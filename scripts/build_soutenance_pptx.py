# -*- coding: utf-8 -*-
"""Soutenance EmoLearn — 29 diapositives, environ 25 minutes."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "rapport" / "figures"
LOGO = ROOT / "rapport" / "logos"
OUT = ROOT / "docs" / "soutenance" / "EmoLearn_soutenance.pptx"

SW, SH = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x0B, 0x2A, 0x4A)
NAVY2 = RGBColor(0x14, 0x3D, 0x66)
INK = RGBColor(0x1C, 0x28, 0x36)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
BLUE = RGBColor(0x1B, 0x4F, 0x8A)
GOLD = RGBColor(0xA6, 0x7C, 0x2D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF4, 0xF7, 0xFB)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE3, 0xE9, 0xF0)
TEAL = RGBColor(0x0F, 0x6E, 0x6A)
SOFT = RGBColor(0xE8, 0xF0, 0xF8)
AMBER = RGBColor(0x8A, 0x4B, 0x12)
RED = RGBColor(0x8C, 0x3A, 0x3A)
GREEN = RGBColor(0x1B, 0x6B, 0x45)

TOTAL = 29


def _no_line(shape):
    shape.line.fill.background()


def rect(slide, l, t, w, h, fill, rounded=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    if rounded:
        sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        _no_line(sh)
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def _set_run(run, text, size, bold, color, italic=False):
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.append(latin)
    latin.set("typeface", "Calibri")


def write(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(line.get("before", 0))
        p.space_after = Pt(line.get("after", 0))
        run = p.add_run()
        _set_run(
            run,
            line["text"],
            line.get("size", 16),
            line.get("bold", False),
            line.get("color", INK),
            line.get("italic", False),
        )
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def chrome(slide, kicker, title, page):
    rect(slide, 0, 0, SW, SH, BG)
    rect(slide, 0, 0, SW, Inches(0.08), NAVY)
    write(
        slide,
        Inches(0.55),
        Inches(0.28),
        Inches(10),
        Inches(0.28),
        [{"text": kicker.upper(), "size": 12, "bold": True, "color": GOLD}],
    )
    write(
        slide,
        Inches(0.55),
        Inches(0.52),
        Inches(12.2),
        Inches(0.7),
        [{"text": title, "size": 28, "bold": True, "color": NAVY}],
    )
    rect(slide, Inches(0.55), Inches(1.22), Inches(1.35), Inches(0.045), GOLD)
    write(
        slide,
        Inches(0.55),
        Inches(7.12),
        Inches(8),
        Inches(0.28),
        [{"text": "EmoLearn  ·  Soutenance de master  ·  Briki Maram", "size": 11, "color": MUTED}],
    )
    write(
        slide,
        Inches(10.4),
        Inches(7.12),
        Inches(2.4),
        Inches(0.28),
        [{"text": f"{page}  /  {TOTAL}", "size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    )


def add_pic(slide, path, l, t, max_w, max_h):
    with Image.open(path) as im:
        aw, ah = im.size
    ratio = aw / ah
    w, h = max_w, max_w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    left = l + (max_w - w) / 2
    top = t + (max_h - h) / 2
    return slide.shapes.add_picture(str(path), left, top, w, h)


def card(slide, l, t, w, h, accent=BLUE):
    rect(slide, l, t, w, h, CARD, rounded=True, line=LINE)
    rect(slide, l, t, Inches(0.08), h, accent)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_title(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(1.35), WHITE)
    rect(s, 0, Inches(1.35), SW, Inches(0.06), GOLD)
    for path, left in (
        (LOGO / "logo_esct.png", Inches(0.45)),
        (LOGO / "logo_isamm.png", Inches(2.15)),
        (LOGO / "logo_uma.png", Inches(11.35)),
    ):
        if path.exists():
            add_pic(s, path, left, Inches(0.18), Inches(1.55), Inches(1.0))
    write(
        s, Inches(3.7), Inches(0.28), Inches(7.4), Inches(0.9),
        [
            {"text": "Université de la Manouba", "size": 14, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER},
            {"text": "École Supérieure de Commerce de Tunis", "size": 13, "color": INK, "align": PP_ALIGN.CENTER, "before": 2},
            {"text": "Master de recherche en informatique décisionnelle", "size": 12, "color": MUTED, "align": PP_ALIGN.CENTER, "before": 4},
        ],
    )
    write(s, Inches(0.8), Inches(1.85), Inches(11.7), Inches(0.35), [{"text": "MÉMOIRE DE MASTER DE RECHERCHE", "size": 13, "bold": True, "color": GOLD, "align": PP_ALIGN.CENTER}])
    write(
        s, Inches(0.9), Inches(2.35), Inches(11.5), Inches(2.0),
        [
            {"text": "EmoLearn : personnalisation pédagogique", "size": 32, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER},
            {"text": "adaptative multimodale", "size": 32, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "before": 2},
        ],
    )
    write(
        s, Inches(1.4), Inches(4.5), Inches(10.5), Inches(0.7),
        [
            {"text": "Détection de l'engagement émotionnel et génération de contenu", "size": 18, "color": RGBColor(0xD6, 0xE2, 0xF0), "align": PP_ALIGN.CENTER},
            {"text": "par IA générative", "size": 18, "color": RGBColor(0xD6, 0xE2, 0xF0), "align": PP_ALIGN.CENTER, "before": 2},
        ],
    )
    rect(s, Inches(5.9), Inches(5.4), Inches(1.5), Inches(0.035), GOLD)
    write(
        s, Inches(0.8), Inches(5.65), Inches(11.7), Inches(1.3),
        [
            {"text": "Briki Maram", "size": 22, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER},
            {"text": "Encadrante académique : Dr REZGUI Kalthom", "size": 15, "color": RGBColor(0xD6, 0xE2, 0xF0), "align": PP_ALIGN.CENTER, "before": 8},
            {"text": "Année universitaire 2025 – 2026", "size": 14, "color": RGBColor(0xB7, 0xC7, 0xD8), "align": PP_ALIGN.CENTER, "before": 6},
        ],
    )
    notes(s, "30 secondes. Saluer, donner le titre.\n« EmoLearn relie la détection de l'engagement et la génération d'un contenu adapté. C'est un prototype de recherche, pas un tuteur déjà en ligne. »")


def slide_plan(prs):
    s = blank(prs)
    chrome(s, "Plan", "Vingt-cinq minutes, six temps", 2)
    items = [
        ("01", "Problème", "Le MOOC, les quatre états, la question."),
        ("02", "Chaîne", "Architecture, données, deux pipelines."),
        ("03", "Perception", "Visage, voix, fusion."),
        ("04", "Pédagogie", "Stratégies, patrons, séquence, LLM, RAG."),
        ("05", "Preuve", "Protocole, chiffres, démonstration."),
        ("06", "Lecture", "Limites, ce qui n'est pas montré, suite."),
    ]
    positions = [
        (Inches(0.55), Inches(1.6)),
        (Inches(6.85), Inches(1.6)),
        (Inches(0.55), Inches(3.35)),
        (Inches(6.85), Inches(3.35)),
        (Inches(0.55), Inches(5.1)),
        (Inches(6.85), Inches(5.1)),
    ]
    for (num, title, body), (x, y) in zip(items, positions):
        card(s, x, y, Inches(5.95), Inches(1.55), BLUE if num in ("01", "03", "05") else TEAL)
        write(s, x + Inches(0.28), y + Inches(0.22), Inches(1.1), Inches(0.4), [{"text": num, "size": 18, "bold": True, "color": GOLD}])
        write(s, x + Inches(1.3), y + Inches(0.24), Inches(4.3), Inches(0.4), [{"text": title, "size": 20, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.28), y + Inches(0.85), Inches(5.4), Inches(0.5), [{"text": body, "size": 15, "color": INK}])
    notes(s, "20 secondes. Nommer les six temps, ne pas les commenter.\n« Je commence par le problème, puis la chaîne, la perception, la décision, les chiffres, et ce que je ne prétends pas. »")


def slide_contexte(prs):
    s = blank(prs)
    chrome(s, "Contexte", "Dans un MOOC, l'enseignant ne voit plus", 3)
    items = [
        ("01", "Le signal disparaît", "Ennui, confusion et frustration précèdent le décrochage. Sans visage ni échange, la plateforme ne régule plus.", BLUE),
        ("02", "Deux avancées séparées", "Les détecteurs d'engagement et les grands modèles de langue progressent. Ils communiquent peu.", TEAL),
        ("03", "Le chaînon manquant", "Ni un classifieur de plus, ni un générateur de plus. Une chaîne : percevoir, décider, générer.", GOLD),
    ]
    x = Inches(0.55)
    for num, title, body, color in items:
        card(s, x, Inches(1.6), Inches(3.9), Inches(3.7), color)
        write(s, x + Inches(0.28), Inches(1.82), Inches(3.4), Inches(0.4), [{"text": num, "size": 18, "bold": True, "color": color}])
        write(s, x + Inches(0.28), Inches(2.4), Inches(3.4), Inches(0.9), [{"text": title, "size": 20, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.28), Inches(3.45), Inches(3.4), Inches(1.55), [{"text": body, "size": 16, "color": INK}])
        x += Inches(4.1)
    rect(s, Inches(0.55), Inches(5.55), Inches(12.23), Inches(1.3), SOFT, rounded=True)
    write(
        s, Inches(0.85), Inches(5.75), Inches(11.7), Inches(0.95),
        [
            {"text": "Ce que je revendique", "size": 14, "bold": True, "color": BLUE},
            {"text": "Un prototype de recherche : la chaîne peut être construite, instrumentée et évaluée. Elle ne réduit pas, ici, le décrochage d'une plateforme réelle.", "size": 16, "color": INK, "before": 4},
        ],
    )
    notes(s, "1 minute.\n« En présentiel, l'enseignant voit l'ennui, la confusion, la frustration. Dans un MOOC, ce signal disparaît. Les détecteurs et les LLM existent, séparément. EmoLearn les relie. Je ne prétends pas baisser le taux d'abandon. »")


def slide_etats(prs):
    s = blank(prs)
    chrome(s, "Engagement", "Quatre états, pas une émotion unique", 4)
    states = [
        ("Ennui", "L'apprenant décroche. Il faut un exemple, ou un exercice.", TEAL),
        ("Engagement", "Il suit. On peut renforcer, ou ne rien envoyer.", BLUE),
        ("Confusion", "Le concept bloque. Il faut reformuler.", GOLD),
        ("Frustration", "Le blocage dure. Il faut soutenir avant d'exiger.", RED),
    ]
    x = Inches(0.55)
    for title, body, color in states:
        card(s, x, Inches(1.6), Inches(2.95), Inches(3.15), color)
        write(s, x + Inches(0.22), Inches(1.85), Inches(2.55), Inches(0.8), [{"text": title, "size": 22, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.22), Inches(2.8), Inches(2.55), Inches(1.6), [{"text": body, "size": 15, "color": INK}])
        x += Inches(3.15)
    rect(s, Inches(0.55), Inches(5.0), Inches(12.23), Inches(1.85), SOFT, rounded=True)
    write(
        s, Inches(0.85), Inches(5.2), Inches(11.7), Inches(1.5),
        [
            {"text": "DAiSEE les note de 0 à 3", "size": 16, "bold": True, "color": NAVY},
            {"text": "L'engagement est concentré sur 2 et 3. La confusion forte est rare. Une accuracy globale est donc trompeuse : je rapporte le F1 macro, et le rappel des états urgents.", "size": 16, "color": INK, "before": 6},
        ],
    )
    notes(s, "45 secondes.\n« Quatre états d'apprentissage, pas les sept émotions du visage. Ennui, engagement, confusion, frustration, de 0 à 3. Le corpus est déséquilibré : beaucoup d'engagement, peu de confusion forte. C'est pour cela que je ne mets pas en avant une accuracy. »")


def slide_position(prs):
    s = blank(prs)
    chrome(s, "Positionnement", "Le mémoire se place entre deux littératures", 5)
    card(s, Inches(0.55), Inches(1.6), Inches(6.0), Inches(3.55), BLUE)
    write(s, Inches(0.85), Inches(1.85), Inches(5.4), Inches(0.4), [{"text": "Perception", "size": 20, "bold": True, "color": NAVY}])
    write(s, Inches(0.85), Inches(2.45), Inches(5.4), Inches(2.3), [{"text": "Les modèles multimodaux estiment l'état affectif d'un apprenant. DAiSEE pour le visage en situation d'apprentissage. IEMOCAP pour la voix.", "size": 16, "color": INK}])
    card(s, Inches(6.78), Inches(1.6), Inches(6.0), Inches(3.55), TEAL)
    write(s, Inches(7.08), Inches(1.85), Inches(5.4), Inches(0.4), [{"text": "Génération", "size": 20, "bold": True, "color": NAVY}])
    write(s, Inches(7.08), Inches(2.45), Inches(5.4), Inches(2.3), [{"text": "Les grands modèles de langue rendent un texte pédagogique fluide. Sans cadre, ils hallucinent, et ils ne savent pas quand se taire.", "size": 16, "color": INK}])
    rect(s, Inches(0.55), Inches(5.4), Inches(12.23), Inches(1.45), NAVY, rounded=True)
    write(
        s, Inches(0.85), Inches(5.6), Inches(11.7), Inches(1.1),
        [{"text": "EmoLearn n'ajoute pas un détecteur, ni un chatbot. Il reconstruit la boucle : percevoir, décider, générer, journaliser.", "size": 18, "color": WHITE}],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(s, "40 secondes.\n« La littérature avance sur deux fronts qui se parlent peu. Moi, je prends les deux et je les mets dans le même logiciel, avec une décision traçable entre les deux. »")


def slide_question(prs):
    s = blank(prs)
    chrome(s, "Question", "La question qui guide le mémoire", 6)
    rect(s, Inches(0.55), Inches(1.7), Inches(12.23), Inches(2.15), NAVY, rounded=True)
    write(
        s, Inches(0.9), Inches(1.95), Inches(11.55), Inches(1.7),
        [{"text": "Comment inférer l'état d'engagement à partir de signaux multimodaux, puis générer un contenu adapté, tout en gardant la décision pédagogique traçable ?", "size": 24, "color": WHITE}],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    write(
        s, Inches(0.55), Inches(4.2), Inches(12.2), Inches(2.2),
        [
            {"text": "La réponse n'est pas un oui.", "size": 22, "bold": True, "color": NAVY},
            {"text": "C'est une réponse de faisabilité : oui, dans un prototype simulé, avec des scores modestes, et une validation humaine encore ouverte.", "size": 18, "color": INK, "before": 12},
            {"text": "Quatre sous-questions découpent ce oui : percevoir, décider, générer, évaluer.", "size": 18, "color": INK, "before": 10},
        ],
    )
    notes(s, "40 secondes. Lire la question une fois.\n« Je ne réponds pas par un oui. Je réponds : c'est faisable en simulation, les scores restent modestes, et les enseignants n'ont pas encore noté l'ensemble. »")


def slide_qr(prs):
    s = blank(prs)
    chrome(s, "Sous-questions", "Quatre questions, quatre verdicts en fin d'oral", 7)
    rows = [
        ("QR1", "Percevoir", "Estimer l'état depuis le visage DAiSEE et la voix IEMOCAP. Que gagne la fusion ?"),
        ("QR2", "Décider", "Relier cet état à une stratégie que l'on peut justifier, et parfois à l'absence d'intervention."),
        ("QR3", "Générer", "Un LLM guidé par vingt patrons, ancré par un RAG, produit-il un contenu adapté ?"),
        ("QR4", "Évaluer", "Quel protocole juge la perception, la décision et le texte, sans masquer les limites ?"),
    ]
    y = Inches(1.6)
    for code, title, body in rows:
        rect(s, Inches(0.55), y, Inches(1.25), Inches(1.15), SOFT, rounded=True)
        write(s, Inches(0.55), y, Inches(1.25), Inches(1.15), [{"text": code, "size": 16, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(2.05), y + Inches(0.15), Inches(10.4), Inches(0.4), [{"text": title, "size": 20, "bold": True, "color": NAVY}])
        write(s, Inches(2.05), y + Inches(0.58), Inches(10.4), Inches(0.45), [{"text": body, "size": 16, "color": INK}])
        y += Inches(1.3)
    notes(s, "40 secondes. Les nommer, ne pas les trancher.\n« Percevoir, décider, générer, évaluer. Le verdict de chacune est sur l'avant-dernière diapositive. »")


def slide_contrib(prs):
    s = blank(prs)
    chrome(s, "Contributions", "Cinq apports, pas un modèle de plus", 8)
    rows = [
        ("C1", "Architecture", "API, interface, clip, séquence, tuteur, journal. La décision est celle du runtime."),
        ("C2", "Adaptation", "Cascade A, B, C, D, vingt patrons, règles temporelles, non-intervention."),
        ("C3", "Perception", "FER+AU et WavLM, avec les ablations qui disent ce qui a été retenu, et pourquoi."),
        ("C4", "Génération", "LLM contraint, RAG, environ 220 contenus, audit de forme."),
        ("C5", "Évaluation", "Quatre niveaux. Les limites sont mesurées, pas cachées dans une annexe."),
    ]
    y = Inches(1.5)
    for code, title, body in rows:
        card(s, Inches(0.55), y, Inches(12.23), Inches(1.0), BLUE)
        write(s, Inches(0.8), y, Inches(0.9), Inches(1.0), [{"text": code, "size": 16, "bold": True, "color": GOLD}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(1.8), y + Inches(0.12), Inches(2.3), Inches(0.75), [{"text": title, "size": 16, "bold": True, "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(4.2), y + Inches(0.12), Inches(8.2), Inches(0.75), [{"text": body, "size": 15, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.08)
    notes(s, "50 secondes.\n« Cinq contributions. L'architecture de bout en bout. Le modèle qui va de l'état à la stratégie. La perception, avec ses ablations. La génération cadrée. Et une évaluation qui dit aussi ce qui ne marche pas. »")


def slide_archi(prs):
    s = blank(prs)
    chrome(s, "Architecture", "Percevoir, décider, générer", 9)
    write(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(0.32), [{"text": "Le code décide quoi produire. Le modèle de langue rédige dans ce cadre.", "size": 15, "color": MUTED}])
    add_pic(s, FIG / "emolearn_architecture.png", Inches(1.3), Inches(1.75), Inches(10.7), Inches(5.15))
    notes(s, "1 minute. Montrer la flèche, ne pas lire les boîtes.\n« Clip DAiSEE, audio IEMOCAP. Visage et voix estiment quatre états. Fusion 60 / 40, pondérée par la confiance. Le moteur choisit un patron, ou ne fait rien. Le RAG ancre, le LLM rédige. Tout est journalisé. »")


def slide_donnees(prs):
    s = blank(prs)
    chrome(s, "Données", "Deux corpus, et un appariement assumé", 10)
    card(s, Inches(0.55), Inches(1.55), Inches(4.0), Inches(3.9), BLUE)
    write(s, Inches(0.85), Inches(1.75), Inches(3.5), Inches(0.4), [{"text": "DAiSEE", "size": 20, "bold": True, "color": NAVY}])
    write(s, Inches(0.85), Inches(2.35), Inches(3.5), Inches(2.7), [{"text": "8 571 clips annotés. Quatre dimensions, de 0 à 3. C'est le visage. Il n'apporte pas un audio d'apprentissage exploitable.", "size": 16, "color": INK}])
    card(s, Inches(4.7), Inches(1.55), Inches(4.0), Inches(3.9), TEAL)
    write(s, Inches(5.0), Inches(1.75), Inches(3.5), Inches(0.4), [{"text": "IEMOCAP", "size": 20, "bold": True, "color": NAVY}])
    write(s, Inches(5.0), Inches(2.35), Inches(3.5), Inches(2.7), [{"text": "La voix réelle. Émotions de base, pas les quatre états DAiSEE. Je projette ces émotions vers ennui, engagement, confusion, frustration.", "size": 16, "color": INK}])
    card(s, Inches(8.85), Inches(1.55), Inches(3.95), Inches(3.9), GOLD)
    write(s, Inches(9.15), Inches(1.75), Inches(3.4), Inches(0.4), [{"text": "Appariement", "size": 20, "bold": True, "color": NAVY}])
    write(s, Inches(9.15), Inches(2.35), Inches(3.4), Inches(2.7), [{"text": "Filtre sur l'état dominant du clip, puis similarité de profil. Ce n'est pas la voix de la personne filmée.", "size": 16, "color": INK}])
    rect(s, Inches(0.55), Inches(5.65), Inches(12.23), Inches(1.2), RGBColor(0xFB, 0xF3, 0xE4), rounded=True)
    write(s, Inches(0.85), Inches(5.85), Inches(11.7), Inches(0.85), [{"text": "Le prototype ne rejoue pas les 8 571 clips. Il est exercé sur 50 scénarios stratifiés, pour couvrir les états rares.", "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "1 minute.\n« DAiSEE, c'est le visage en situation d'apprentissage, sans audio utile. IEMOCAP, c'est la voix, avec d'autres étiquettes. Je les apparie. Je le dis clairement : ce n'est pas la même personne. Les 50 scénarios sont un échantillon stratifié, pas tout le corpus. »")


def slide_visage(prs):
    s = blank(prs)
    chrome(s, "Visage", "FER + unités d'action, le juge retenu", 11)
    write(s, Inches(0.55), Inches(1.55), Inches(7.2), Inches(0.7), [{"text": "F1 macro  0,37", "size": 40, "bold": True, "color": NAVY}])
    write(s, Inches(0.55), Inches(2.4), Inches(7.2), Inches(0.4), [{"text": "Heuristique, même protocole : 0,19", "size": 18, "color": MUTED}])
    steps = [
        ("1", "ViT pré-entraîné : 7 émotions du visage."),
        ("2", "MediaPipe : 10 unités d'action, proxy des muscles."),
        ("3", "Forêt aléatoire : ces 17 nombres vers 4 états DAiSEE."),
    ]
    y = Inches(3.05)
    for num, text in steps:
        rect(s, Inches(0.55), y, Inches(0.42), Inches(0.42), NAVY, rounded=True)
        write(s, Inches(0.55), y, Inches(0.42), Inches(0.42), [{"text": num, "size": 14, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(1.15), y, Inches(6.3), Inches(0.42), [{"text": text, "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.62)
    card(s, Inches(8.15), Inches(1.55), Inches(4.65), Inches(5.2), GOLD)
    write(s, Inches(8.45), Inches(1.8), Inches(4.15), Inches(1.3), [{"text": "Rappel des états urgents", "size": 16, "bold": True, "color": AMBER}])
    write(s, Inches(8.45), Inches(3.0), Inches(4.15), Inches(0.8), [{"text": "34 %", "size": 40, "bold": True, "color": NAVY}])
    write(s, Inches(8.45), Inches(4.0), Inches(4.15), Inches(2.2), [{"text": "Frustration et confusion, chacune. Le ViT fine-tuné filtre les frames hésitantes. Il ne choisit pas l'état envoyé à la fusion.", "size": 16, "color": INK}])
    notes(s, "1 minute.\n« Le juge, c'est FER plus unités d'action. Sept émotions, dix mouvements du visage, une forêt aléatoire. F1 0,37 contre 0,19. Le rappel de la frustration et de la confusion est de 34 %. Le ViT que j'ai fine-tuné sert de filtre, pas de décision. »")


def slide_voix(prs):
    s = blank(prs)
    chrome(s, "Voix", "WavLM, tête seule, retenu", 12)
    card(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(3.7), TEAL)
    write(s, Inches(0.85), Inches(1.75), Inches(5.4), Inches(0.35), [{"text": "RETENU", "size": 13, "bold": True, "color": TEAL}])
    write(s, Inches(0.85), Inches(2.2), Inches(5.4), Inches(0.7), [{"text": "F1  0,43", "size": 36, "bold": True, "color": NAVY}])
    write(s, Inches(0.85), Inches(3.15), Inches(5.4), Inches(1.7), [{"text": "WavLM, tête seule, vers les quatre états DAiSEE. Un second modèle, ancré sur les émotions IEMOCAP, intervient surtout si le premier hésite sur la confusion.", "size": 16, "color": INK}])
    card(s, Inches(6.78), Inches(1.55), Inches(6.0), Inches(3.7), RED)
    write(s, Inches(7.08), Inches(1.75), Inches(5.4), Inches(0.35), [{"text": "ESSAYÉ, NON RETENU", "size": 13, "bold": True, "color": RED}])
    write(s, Inches(7.08), Inches(2.2), Inches(5.4), Inches(0.7), [{"text": "F1  0,40", "size": 36, "bold": True, "color": NAVY}])
    write(s, Inches(7.08), Inches(3.15), Inches(5.4), Inches(1.7), [{"text": "Le même WavLM, réseau entier dégelé. L'accuracy est proche. Le F1 baisse. Je garde la tête seule.", "size": 16, "color": INK}])
    rect(s, Inches(0.55), Inches(5.5), Inches(12.23), Inches(1.35), RGBColor(0xFB, 0xF3, 0xE4), rounded=True)
    write(s, Inches(0.85), Inches(5.7), Inches(11.7), Inches(1.0), [{"text": "Cet audio est un énoncé IEMOCAP apparié au clip. Ce n'est pas la voix de l'apprenant filmé dans DAiSEE.", "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "50 secondes.\n« WavLM, tête seule, F1 0,43. J'ai dégelé tout le réseau : le F1 tombe à 0,40, je ne le retiens pas. L'audio n'est pas celui de la vidéo. »")


def slide_fusion(prs):
    s = blank(prs)
    chrome(s, "Fusion", "60 % visage, 40 % voix, puis la confiance", 13)
    items = [
        ("60 / 40", "Prior", "Le visage pèse plus que la voix. Ce rapport vient d'un balayage, pas d'un choix arbitraire.", BLUE),
        ("× confiance", "Poids effectif", "Chaque modalité est multipliée par sa confiance, puis les poids sont renormalisés.", TEAL),
        ("Repli", "Si tout est nul", "Quand les deux confiances sont nulles, une priorité pédagogique décide à leur place.", GOLD),
    ]
    x = Inches(0.55)
    for value, title, body, color in items:
        card(s, x, Inches(1.6), Inches(3.9), Inches(3.85), color)
        write(s, x + Inches(0.28), Inches(1.85), Inches(3.4), Inches(0.7), [{"text": value, "size": 28, "bold": True, "color": color}])
        write(s, x + Inches(0.28), Inches(2.7), Inches(3.4), Inches(0.45), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.28), Inches(3.3), Inches(3.4), Inches(1.8), [{"text": body, "size": 16, "color": INK}])
        x += Inches(4.1)
    rect(s, Inches(0.55), Inches(5.7), Inches(12.23), Inches(1.15), SOFT, rounded=True)
    write(s, Inches(0.85), Inches(5.9), Inches(11.7), Inches(0.8), [{"text": "Les frames de visage trop hésitantes sont écartées avant la moyenne. Seuil 0,35.", "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "45 secondes.\n« Ce n'est pas une moyenne fixe. 60 % visage, 40 % voix, puis chaque côté est pesé par sa confiance. Si les deux sont muets, une règle pédagogique prend le relais. »")


def slide_pipelines(prs):
    s = blank(prs)
    chrome(s, "Deux cadres", "Le même moteur, pas la même vérité", 14)
    card(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(4.35), BLUE)
    write(s, Inches(0.85), Inches(1.75), Inches(5.4), Inches(0.4), [{"text": "Pipeline produit", "size": 20, "bold": True, "color": NAVY}])
    write(
        s, Inches(0.85), Inches(2.4), Inches(5.4), Inches(3.1),
        [
            {"text": "Visage : clip DAiSEE réel.", "size": 16, "color": INK},
            {"text": "Voix : IEMOCAP apparié.", "size": 16, "color": INK, "before": 8},
            {"text": "Vérité : l'état dominant du clip. C'est un proxy.", "size": 16, "color": INK, "before": 8},
            {"text": "On mesure le runtime tel qu'il est servi.", "size": 16, "color": INK, "before": 8},
        ],
    )
    card(s, Inches(6.78), Inches(1.55), Inches(6.0), Inches(4.35), TEAL)
    write(s, Inches(7.08), Inches(1.75), Inches(5.4), Inches(0.4), [{"text": "Banc aligné", "size": 20, "bold": True, "color": NAVY}])
    write(
        s, Inches(7.08), Inches(2.4), Inches(5.4), Inches(3.1),
        [
            {"text": "Visage et voix : même locuteur IEMOCAP.", "size": 16, "color": INK},
            {"text": "Même segment, pas un collage.", "size": 16, "color": INK, "before": 8},
            {"text": "n = 80, rejoué avec FER+AU et WavLM.", "size": 16, "color": INK, "before": 8},
            {"text": "On mesure la fusion, sans le biais d'appariement.", "size": 16, "color": INK, "before": 8},
        ],
    )
    write(s, Inches(0.55), Inches(6.1), Inches(12.2), Inches(0.7), [{"text": "Les deux partagent le moteur. Un score du banc n'est pas un score du produit.", "size": 16, "bold": True, "color": NAVY}])
    notes(s, "50 secondes.\n« Deux pipelines, un seul moteur de fusion. À gauche, ce que l'application sert : visage DAiSEE, voix appariée, pas de vérité conjointe. À droite, un banc où c'est la même personne, pour juger la fusion elle-même. Je ne mélange pas les deux chiffres. »")


def slide_abcd(prs):
    s = blank(prs)
    chrome(s, "Décision", "L'état choisit la stratégie", 15)
    strategies = [
        ("A", "Expliquer", "Confusion : reformuler le concept, avant d'ajouter un exercice.", BLUE),
        ("B", "Concrétiser", "Ennui ou engagement faible : un exemple pris dans un geste familier.", TEAL),
        ("C", "Exercer", "Ennui élevé : un exercice plus exigeant, pas une redite.", GOLD),
        ("D", "Soutenir", "Frustration : rassurer et simplifier. Engagement haut : renforcer.", RED),
    ]
    x = Inches(0.55)
    for code, title, body, color in strategies:
        card(s, x, Inches(1.6), Inches(2.95), Inches(3.55), color)
        write(s, x + Inches(0.22), Inches(1.8), Inches(2.55), Inches(0.55), [{"text": code, "size": 28, "bold": True, "color": color}])
        write(s, x + Inches(0.22), Inches(2.5), Inches(2.55), Inches(0.45), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.22), Inches(3.1), Inches(2.55), Inches(1.7), [{"text": body, "size": 15, "color": INK}])
        x += Inches(3.15)
    rect(s, Inches(0.55), Inches(5.4), Inches(12.23), Inches(1.45), NAVY, rounded=True)
    write(s, Inches(0.85), Inches(5.65), Inches(11.7), Inches(1.0), [{"text": "Le modèle de langue ne choisit pas A, B, C ou D. Il reçoit le patron déjà décidé, et il rédige.", "size": 18, "color": WHITE}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "45 secondes.\n« Quatre familles. A, on explique. B, on donne un exemple. C, on fait exercer. D, on soutient, ou on renforce. Le LLM ne choisit pas la lettre. Le moteur l'a déjà choisie. »")


def slide_patrons(prs):
    s = blank(prs)
    chrome(s, "Patrons", "Vingt façons d'intervenir, et une de se taire", 16)
    write(s, Inches(0.55), Inches(1.42), Inches(12.2), Inches(0.4), [{"text": "A1 à D5. Le patron fixe le geste pédagogique. Le texte change, la stratégie non.", "size": 16, "color": INK}])
    add_pic(s, FIG / "emolearn_decision_pipeline.png", Inches(0.35), Inches(1.9), Inches(12.6), Inches(2.7))
    examples = [
        ("A3", "Déblocage", "Confusion et frustration ensemble."),
        ("B2", "Exemple", "Un cas concret, court."),
        ("C2", "Exercice", "Une tâche progressive."),
        ("Aucune", "Silence", "Engagement déjà stable."),
    ]
    x = Inches(0.55)
    for code, title, body in examples:
        card(s, x, Inches(4.8), Inches(2.95), Inches(2.0), BLUE if code != "Aucune" else TEAL)
        write(s, x + Inches(0.2), Inches(4.95), Inches(2.55), Inches(0.35), [{"text": code, "size": 14, "bold": True, "color": GOLD}])
        write(s, x + Inches(0.2), Inches(5.3), Inches(2.55), Inches(0.35), [{"text": title, "size": 16, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.2), Inches(5.7), Inches(2.55), Inches(0.85), [{"text": body, "size": 13, "color": INK}])
        x += Inches(3.15)
    notes(s, "40 secondes.\n« Vingt patrons. A3 débloque, B2 donne un exemple, C2 fait exercer. Et « aucune » est une sortie du moteur : si l'apprenant est engagé et stable, ne rien envoyer est la bonne décision. »")


def slide_sequence(prs):
    s = blank(prs)
    chrome(s, "Séquence", "Cinq clips, le même apprenant", 17)
    write(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(0.7), [{"text": "Le mode clip est un instant. Le mode séquence est une mini-séance : cinq extraits du même visage DAiSEE, analysés un par un.", "size": 16, "color": INK}])
    rules = [
        ("Immédiate", "Confusion très haute et frustration, ou engagement à zéro, sur le clip courant.", RED),
        ("Rapide", "Ennui, confusion ou frustration élevés sur deux clips de suite.", GOLD),
        ("Différée", "Engagement faible, stable, sur quatre clips.", BLUE),
        ("Aucune", "Engagement haut, stable sur deux clips : pas d'intervention.", TEAL),
    ]
    y = Inches(2.3)
    for title, body, color in rules:
        card(s, Inches(0.55), y, Inches(12.23), Inches(1.05), color)
        write(s, Inches(0.85), y, Inches(2.4), Inches(1.05), [{"text": title, "size": 16, "bold": True, "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(3.4), y, Inches(9.0), Inches(1.05), [{"text": body, "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.15)
    notes(s, "45 secondes.\n« Cinq extraits du même apprenant, pas un collage de visages. Chaque extrait repasse par le pipeline. Si la difficulté dure, l'urgence monte : immédiate, rapide, différée. Si l'engagement tient, on se tait. La voix, elle, est ré-appariée à chaque clip : ce n'est toujours pas la sienne. »")


def slide_llm(prs):
    s = blank(prs)
    chrome(s, "Génération", "Le modèle rédige. Il ne décide pas.", 18)
    facts = [
        ("Groq", "GPT-OSS 20B, appelé par API. Le modèle n'est pas fine-tuné sur ce mémoire."),
        ("0,3", "Température basse : moins de digressions, une réponse pédagogique courte."),
        ("Patron", "Le prompt impose la stratégie déjà choisie, le niveau, et le contexte du cours."),
        ("Tuteur", "Le chat reprend l'état. Sa température change : plus basse si l'apprenant est confus."),
    ]
    positions = [(Inches(0.55), Inches(1.6)), (Inches(6.85), Inches(1.6)), (Inches(0.55), Inches(4.15)), (Inches(6.85), Inches(4.15))]
    for (title, body), (x, y) in zip(facts, positions):
        card(s, x, y, Inches(5.95), Inches(2.3), BLUE)
        write(s, x + Inches(0.3), y + Inches(0.25), Inches(5.4), Inches(0.5), [{"text": title, "size": 22, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.3), y + Inches(0.9), Inches(5.4), Inches(1.1), [{"text": body, "size": 16, "color": INK}])
    notes(s, "40 secondes.\n« Groq, GPT-OSS 20 milliards, température 0,3. Je ne l'ai pas ré-entraîné. Le patron est déjà dans le prompt. Le tuteur conversationnel existe : il parle plus prudemment quand l'état est la confusion. »")


def slide_rag(prs):
    s = blank(prs)
    chrome(s, "RAG", "Vingt fiches, trois documents, un seuil", 19)
    metrics = [
        ("20", "fiches", "La base, pas un cours entier."),
        ("3", "documents", "Injectés dans le prompt."),
        ("0,3", "seuil", "En dessous, le document est écarté."),
    ]
    x = Inches(0.55)
    for value, title, body in metrics:
        card(s, x, Inches(1.6), Inches(3.9), Inches(2.7), TEAL)
        write(s, x + Inches(0.28), Inches(1.8), Inches(3.4), Inches(0.7), [{"text": value, "size": 36, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.28), Inches(2.6), Inches(3.4), Inches(0.4), [{"text": title, "size": 18, "bold": True, "color": TEAL}])
        write(s, x + Inches(0.28), Inches(3.15), Inches(3.4), Inches(0.8), [{"text": body, "size": 16, "color": INK}])
        x += Inches(4.1)
    rect(s, Inches(0.55), Inches(4.6), Inches(12.23), Inches(2.2), SOFT, rounded=True)
    write(
        s, Inches(0.85), Inches(4.85), Inches(11.7), Inches(1.75),
        [
            {"text": "Ce que le RAG fait, et ce qu'il ne fait pas", "size": 16, "bold": True, "color": NAVY},
            {"text": "Il ancre le texte dans des documents vérifiés. Il réduit les inventions. Il ne les supprime pas : le mémoire montre une hallucination chiffrée qui est passée.", "size": 16, "color": INK, "before": 8},
            {"text": "Les embeddings sont multilingues. La recherche est sémantique, pas un mot-clé.", "size": 16, "color": INK, "before": 8},
        ],
    )
    notes(s, "40 secondes.\n« Vingt fiches, les trois plus proches, seuil 0,3. Le RAG diminue les hallucinations. Une erreur chiffrée est quand même dans l'annexe : je ne dis pas que le problème est réglé. »")


def slide_corpus(prs):
    s = blank(prs)
    chrome(s, "Corpus", "La forme tient. Le jugement enseignant est mince.", 20)
    stats = [
        ("~220", "contenus générés"),
        ("97,7 %", "passent l'audit de forme"),
        ("222", "mots en moyenne"),
        ("3,9 / 5", "avis enseignant, n = 3"),
    ]
    x = Inches(0.55)
    for value, label in stats:
        card(s, x, Inches(1.6), Inches(2.95), Inches(2.35), BLUE)
        write(s, x + Inches(0.18), Inches(1.8), Inches(2.6), Inches(0.8), [{"text": value, "size": 26, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.18), Inches(2.75), Inches(2.6), Inches(0.9), [{"text": label, "size": 15, "color": INK}])
        x += Inches(3.15)
    rect(s, Inches(0.55), Inches(4.2), Inches(12.23), Inches(2.55), RGBColor(0xFB, 0xF3, 0xE4), rounded=True)
    write(
        s, Inches(0.85), Inches(4.45), Inches(11.7), Inches(2.1),
        [
            {"text": "Comment lire ces chiffres", "size": 16, "bold": True, "color": AMBER},
            {"text": "BLEU et ROUGE comparent les textes d'un même patron entre eux. Ce n'est pas une copie d'enseignant.", "size": 16, "color": INK, "before": 8},
            {"text": "L'avis 3,9 sur 5 porte sur 8 extraits, lus par 3 personnes. Pas sur les 58 fiches prévues.", "size": 16, "color": INK, "before": 8},
        ],
    )
    notes(s, "45 secondes.\n« Environ 220 textes, 97,7 % conformes à l'audit, autour de 220 mots. BLEU et ROUGE disent que les textes d'un même patron se ressemblent. Trois enseignants ont noté huit extraits : 3,9 sur 5. C'est un signal, pas une validation. »")


def slide_demo(prs):
    s = blank(prs)
    chrome(s, "Démonstration", "Un clip, puis une séquence", 21)
    steps = [
        ("1", "Ouvrir un clip en frustration ou en confusion."),
        ("2", "Montrer le désaccord possible entre visage et voix."),
        ("3", "La fusion retient un état, le moteur choisit le patron."),
        ("4", "Passer en séquence : cinq clips, l'urgence qui change."),
    ]
    y = Inches(1.55)
    for num, text in steps:
        rect(s, Inches(0.55), y, Inches(0.5), Inches(0.5), NAVY, rounded=True)
        write(s, Inches(0.55), y, Inches(0.5), Inches(0.5), [{"text": num, "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(1.25), y + Inches(0.05), Inches(5.3), Inches(0.45), [{"text": text, "size": 15, "color": INK}])
        y += Inches(0.72)
    rect(s, Inches(0.55), Inches(4.55), Inches(6.05), Inches(2.25), RGBColor(0xFB, 0xF3, 0xE4), rounded=True)
    write(
        s, Inches(0.8), Inches(4.75), Inches(5.6), Inches(1.9),
        [
            {"text": "Ce que montre la capture", "size": 14, "bold": True, "color": AMBER},
            {"text": "Frustration retenue, patron B2, exemple concret, sources RAG en bas. Le LLM a rédigé. Il n'a pas choisi B2.", "size": 15, "color": INK, "before": 6},
        ],
    )
    rect(s, Inches(6.85), Inches(1.55), Inches(5.95), Inches(5.25), CARD, rounded=True, line=LINE)
    add_pic(s, FIG / "ui_emolearn_resultats.png", Inches(7.0), Inches(1.7), Inches(5.65), Inches(4.95))
    notes(s, "2 à 3 minutes, sur l'outil si l'API répond.\nDire : le visage et la voix peuvent diverger, la fusion tranche, le patron est choisi avant la rédaction. Puis la séquence. Si l'API tombe, rester sur la capture : B2, frustration, sources visibles. Ne pas déboguer.")


def slide_protocole(prs):
    s = blank(prs)
    chrome(s, "Protocole", "Quatre niveaux, pas un seul score", 22)
    levels = [
        ("1", "Perception", "F1 du visage, F1 de la voix, banc aligné pour la fusion.", "Fait"),
        ("2", "Décision", "49 scénarios : famille pédagogique, état fusionné, type admissible.", "Fait"),
        ("3", "Contenu", "Audit automatique sur le corpus. Avis enseignant sur 8 textes.", "Partiel"),
        ("4", "Usage", "Enquête d'idée, 15 personnes. Le pilote sur le prototype n'a pas eu lieu.", "Ouvert"),
    ]
    y = Inches(1.55)
    for num, title, body, tag in levels:
        card(s, Inches(0.55), y, Inches(12.23), Inches(1.22), BLUE)
        write(s, Inches(0.8), y, Inches(0.7), Inches(1.22), [{"text": num, "size": 22, "bold": True, "color": GOLD}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(1.6), y + Inches(0.14), Inches(2.4), Inches(0.95), [{"text": title, "size": 18, "bold": True, "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(4.1), y + Inches(0.14), Inches(6.3), Inches(0.95), [{"text": body, "size": 15, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(10.5), y, Inches(2.0), Inches(1.22), [{"text": tag, "size": 14, "bold": True, "color": TEAL if tag == "Fait" else AMBER, "align": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.32)
    notes(s, "40 secondes.\n« Quatre niveaux. La perception est mesurée. La décision aussi, sur 49 scénarios. Le contenu est audité, et à peine noté par des enseignants. L'usage réel n'est pas testé : quinze personnes ont jugé l'idée, pas le prototype. »")


def slide_res_perc(prs):
    s = blank(prs)
    chrome(s, "Résultats", "Perception et fusion", 23)
    metrics = [
        ("0,37", "Visage", "F1 macro FER+AU", "Heuristique : 0,19", BLUE),
        ("0,43", "Voix", "F1 macro WavLM", "Réseau entier : 0,40", TEAL),
        ("0,49", "Banc aligné", "Accuracy, n = 80", "Même locuteur", GOLD),
        ("0,35", "États urgents", "Rappel sur ce banc", "Confusion ou frustration", RED),
    ]
    x = Inches(0.55)
    for value, title, subtitle, note, color in metrics:
        card(s, x, Inches(1.6), Inches(2.95), Inches(3.7), color)
        write(s, x + Inches(0.18), Inches(1.85), Inches(2.6), Inches(0.85), [{"text": value, "size": 36, "bold": True, "color": color}])
        write(s, x + Inches(0.18), Inches(2.85), Inches(2.6), Inches(0.4), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.18), Inches(3.35), Inches(2.6), Inches(0.8), [{"text": subtitle, "size": 14, "color": INK}])
        write(s, x + Inches(0.18), Inches(4.3), Inches(2.6), Inches(0.7), [{"text": note, "size": 14, "bold": True, "color": MUTED}])
        x += Inches(3.15)
    rect(s, Inches(0.55), Inches(5.55), Inches(12.23), Inches(1.28), SOFT, rounded=True)
    write(s, Inches(0.85), Inches(5.75), Inches(11.7), Inches(0.95), [{"text": "L'accuracy 0,72 d'un premier classifieur facial, sur un corpus déséquilibré, n'est pas le résultat. Le résultat est le F1.", "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "1 minute. Un chiffre, une phrase.\n« 0,37 le visage. 0,43 la voix. 0,49 la fusion quand c'est la même personne. 0,35 le rappel des états urgents sur ce banc. J'écarte le 0,72 : c'était une accuracy gonflée par l'engagement majoritaire. »")


def slide_res_dec(prs):
    s = blank(prs)
    chrome(s, "Résultats", "Décision et contenu", 24)
    metrics = [
        ("100 %", "Famille", "A, B, C ou D juste", "Sur 49 scénarios", TEAL),
        ("47 %", "État fusionné", "Concordant avec DAiSEE", "Le proxy, pas un or conjoint", BLUE),
        ("55 %", "Type admissible", "Le patron colle à l'état", "Le reste est à discuter", GOLD),
        ("97,7 %", "Audit", "Textes conformes", "Forme, pas fond enseignant", NAVY),
    ]
    x = Inches(0.55)
    for value, title, subtitle, note, color in metrics:
        card(s, x, Inches(1.6), Inches(2.95), Inches(3.7), color)
        write(s, x + Inches(0.15), Inches(1.85), Inches(2.65), Inches(0.85), [{"text": value, "size": 32, "bold": True, "color": color}])
        write(s, x + Inches(0.18), Inches(2.85), Inches(2.6), Inches(0.45), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.18), Inches(3.4), Inches(2.6), Inches(0.8), [{"text": subtitle, "size": 15, "color": INK}])
        write(s, x + Inches(0.18), Inches(4.35), Inches(2.6), Inches(0.7), [{"text": note, "size": 14, "color": MUTED}])
        x += Inches(3.15)
    rect(s, Inches(0.55), Inches(5.55), Inches(12.23), Inches(1.28), SOFT, rounded=True)
    write(s, Inches(0.85), Inches(5.75), Inches(11.7), Inches(0.95), [{"text": "La stratégie tombe dans la bonne famille. L'état fin, lui, ne tombe juste qu'une fois sur deux.", "size": 16, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "50 secondes.\n« Sur 49 scénarios, la famille A, B, C ou D est toujours la bonne. L'état fusionné ne l'est que dans 47 % des cas. Le type de patron est admissible dans 55 %. Le texte, lui, est régulier : 97,7 % d'audit. Régulier ne veut pas dire validé par un enseignant. »")


def slide_limites(prs):
    s = blank(prs)
    chrome(s, "Limites", "Ce qui conditionne la lecture", 25)
    items = [
        ("01", "Pas de salle", "Aucun apprenant réel, aucune captation en direct, aucune mesure d'apprentissage."),
        ("02", "Signaux appariés", "Hors du banc, le visage et la voix ne viennent pas de la même personne."),
        ("03", "Humains ouverts", "Le panel sur les patrons n'est pas noté. Le pilote sur le prototype non plus."),
        ("04", "Scores modestes", "Un F1 de 0,37 ne pilote pas un MOOC. Le RAG tient en vingt fiches."),
    ]
    positions = [(Inches(0.55), Inches(1.55)), (Inches(6.8), Inches(1.55)), (Inches(0.55), Inches(4.15)), (Inches(6.8), Inches(4.15))]
    for (num, title, body), (x, y) in zip(items, positions):
        card(s, x, y, Inches(6.0), Inches(2.35), RED)
        write(s, x + Inches(0.28), y + Inches(0.22), Inches(1.0), Inches(0.35), [{"text": num, "size": 14, "bold": True, "color": RED}])
        write(s, x + Inches(1.2), y + Inches(0.18), Inches(4.5), Inches(0.4), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.28), y + Inches(0.85), Inches(5.4), Inches(1.2), [{"text": body, "size": 15, "color": INK}])
    notes(s, "40 secondes.\n« Quatre limites, celles du chapitre 6. Pas d'apprenant réel. Visage et voix appariés. La validation enseignante n'est pas faite. Et les scores sont trop modestes pour un déploiement. »")


def slide_pas(prs):
    s = blank(prs)
    chrome(s, "Périmètre", "Ce que ce mémoire n'établit pas", 26)
    rows = [
        ("01", "Pas de victoire multimodale en production", "Il n'existe pas de vérité visage-voix conjointe sur le pipeline servi."),
        ("02", "Pas un gain d'apprentissage", "Les métriques de texte sont de forme et de régularité."),
        ("03", "Pas un effet sur le décrochage", "Aucune cohorte MOOC n'a utilisé EmoLearn."),
        ("04", "Pas la même personne, hors banc", "L'alignement visage-voix n'est démontré que sur les 80 segments du banc."),
    ]
    y = Inches(1.6)
    for num, title, body in rows:
        card(s, Inches(0.55), y, Inches(12.23), Inches(1.2), NAVY)
        write(s, Inches(0.8), y, Inches(0.8), Inches(1.2), [{"text": num, "size": 16, "bold": True, "color": GOLD}], anchor=MSO_ANCHOR.MIDDLE)
        write(s, Inches(1.8), y + Inches(0.15), Inches(10.5), Inches(0.4), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, Inches(1.8), y + Inches(0.58), Inches(10.5), Inches(0.45), [{"text": body, "size": 15, "color": INK}])
        y += Inches(1.3)
    notes(s, "30 secondes. Les dire calmement.\n« Quatre choses que je n'établis pas. Que la fusion bat une modalité seule en production. Que le texte fait mieux apprendre. Que le décrochage baisse. Et que visage et voix sont la même personne en dehors du banc. »")


def slide_perspectives(prs):
    s = blank(prs)
    chrome(s, "Perspectives", "D'abord la preuve, ensuite le modèle", 27)
    items = [
        ("Court terme", "Faire noter les patrons par trois experts, et étendre l'avis enseignant au-delà de huit textes. Pas un nouvel entraînement avant cela.", BLUE),
        ("Moyen terme", "Un corpus où visage et voix sont vraiment joints. Un RAG branché sur le cours cible, pas sur vingt fiches.", TEAL),
        ("Plus tard", "Un pilote : contenu statique contre contenu EmoLearn, environ 30 personnes. La caméra en direct seulement avec un cadre éthique.", GOLD),
    ]
    y = Inches(1.6)
    for title, body, color in items:
        card(s, Inches(0.55), y, Inches(12.23), Inches(1.6), color)
        write(s, Inches(0.9), y + Inches(0.2), Inches(11.5), Inches(0.4), [{"text": title, "size": 18, "bold": True, "color": NAVY}])
        write(s, Inches(0.9), y + Inches(0.7), Inches(11.5), Inches(0.7), [{"text": body, "size": 16, "color": INK}])
        y += Inches(1.75)
    notes(s, "40 secondes.\n« La suite n'est pas un plus gros modèle. C'est d'abord l'avis des enseignants sur les patrons. Ensuite un corpus vraiment aligné, et le matériau du cours dans le RAG. Le pilote avec des étudiants vient après, pas avant. »")


def slide_conclusion(prs):
    s = blank(prs)
    chrome(s, "Conclusion", "Une chaîne évaluable", 28)
    rect(s, Inches(0.55), Inches(1.5), Inches(12.23), Inches(1.7), NAVY, rounded=True)
    write(
        s, Inches(0.85), Inches(1.65), Inches(11.65), Inches(1.4),
        [{"text": "La contribution, c'est la chaîne. Les scores restent modestes, la validation humaine n'est pas terminée. Je ne prétends pas réduire le décrochage.", "size": 22, "color": WHITE}],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    verdicts = [
        ("QR1", "Percevoir", "Partiel", "Gain documenté. Pas d'or multimodal en production."),
        ("QR2", "Décider", "Oui, en ingénierie", "La justification enseignante reste ouverte."),
        ("QR3", "Générer", "Partiel", "La forme tient. Le fond est à peine noté."),
        ("QR4", "Évaluer", "Oui, comme cadre", "Le protocole existe. La campagne humaine non."),
    ]
    x = Inches(0.55)
    for code, title, verdict, body in verdicts:
        card(s, x, Inches(3.45), Inches(2.95), Inches(3.35), BLUE)
        write(s, x + Inches(0.18), Inches(3.6), Inches(2.6), Inches(0.3), [{"text": code, "size": 12, "bold": True, "color": GOLD}])
        write(s, x + Inches(0.18), Inches(3.95), Inches(2.6), Inches(0.35), [{"text": title, "size": 16, "bold": True, "color": NAVY}])
        write(s, x + Inches(0.18), Inches(4.4), Inches(2.6), Inches(0.7), [{"text": verdict, "size": 15, "bold": True, "color": BLUE}])
        write(s, x + Inches(0.18), Inches(5.2), Inches(2.6), Inches(1.3), [{"text": body, "size": 14, "color": INK}])
        x += Inches(3.15)
    notes(s, "30 secondes. Lire la phrase, s'arrêter.\n« La contribution, c'est une chaîne évaluable. Les scores sont modestes, la validation humaine n'est pas terminée, et je ne prétends pas réduire le décrochage. »")


def slide_merci(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, Inches(0.12), SH, GOLD)
    write(s, Inches(0.9), Inches(2.15), Inches(11.5), Inches(1.0), [{"text": "Merci", "size": 54, "bold": True, "color": WHITE}])
    write(s, Inches(0.9), Inches(3.3), Inches(11), Inches(0.6), [{"text": "Je vous écoute.", "size": 24, "color": RGBColor(0xD6, 0xE2, 0xF0)}])
    rect(s, Inches(0.9), Inches(4.15), Inches(1.5), Inches(0.04), GOLD)
    write(
        s, Inches(0.9), Inches(4.45), Inches(11), Inches(1.4),
        [
            {"text": "Briki Maram", "size": 18, "bold": True, "color": WHITE},
            {"text": "Encadrante : Dr REZGUI Kalthom  ·  ESCT", "size": 16, "color": RGBColor(0xD6, 0xE2, 0xF0), "before": 6},
            {"text": "EmoLearn  ·  2025 – 2026", "size": 16, "color": RGBColor(0xB7, 0xC7, 0xD8), "before": 4},
        ],
    )
    notes(s, "S'arrêter.\nRéflexes : prototype simulé. Appariement assumé. F1, pas le 0,72. Patrons non notés par un panel. Ne pas ouvrir un notebook sauf si on le demande.")


def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    prs.core_properties.title = "EmoLearn — Soutenance de master"
    prs.core_properties.author = "Briki Maram"
    prs.core_properties.subject = "Personnalisation pédagogique adaptative multimodale"
    builders = [
        slide_title, slide_plan, slide_contexte, slide_etats, slide_position,
        slide_question, slide_qr, slide_contrib, slide_archi, slide_donnees,
        slide_visage, slide_voix, slide_fusion, slide_pipelines, slide_abcd,
        slide_patrons, slide_sequence, slide_llm, slide_rag, slide_corpus,
        slide_demo, slide_protocole, slide_res_perc, slide_res_dec, slide_limites,
        slide_pas, slide_perspectives, slide_conclusion, slide_merci,
    ]
    assert len(builders) == TOTAL
    for build in builders:
        build(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(OUT)
        print(OUT)
    except PermissionError:
        alt = OUT.with_name("EmoLearn_soutenance_29.pptx")
        prs.save(alt)
        print("LOCKED", alt)


if __name__ == "__main__":
    main()
